"""preprocess_opt v3 测试：变更记录页签格式（2026-10-14 与业务侧确认）。

场景构造对齐 mapping-format.md「变更记录页签」规范：三页签 mapping（实体级/属性级/
变更记录）+ RS（口径材料）。批次过滤/改名合成/类型归类/conflict/dangling/漏标/e2e。
不连库。
"""
import json
from pathlib import Path

from openpyxl import Workbook

import pytest

from assemble_ts_baseline import build_ts_baseline
from preprocess_opt import (
    read_change_sheet, pick_batch_version, synth_renames, extract_and_check,
    extract_version_section, parse_change_log, norm_asset, main,
)

PKG_XLSX = "dwb_trade_order_d_全量mapping.xlsx"
PKG_RS = "RS_需求文档.md"

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "opt"

ENTITY_COLS = ["源表schema", "源表物理表名", "源表别名", "目标表逻辑schema",
               "目标表物理名称", "关联&限定条件", "备注", "源表中文名", "目标表中文名"]
ATTR_COLS = ["分组", "源表别名", "源表物理表名", "源表字段名", "映射规则", "映射表达式",
             "目标字段名", "目标字段中文名", "目标字段类型", "备注", "源schema", "源表字段类型"]
CHANGE_COLS = ["变更序号", "变更时间", "变更页签", "变更KEY", "变更列", "变更类型", "变更前", "变更后"]  # 真实模板形态（KEY 大写）

RS_MD = """# 某资产需求规格

## 4. 字段需求

202610版本新增需求：新增 channel_name（渠道名称），来源渠道维表，按订单关联取渠道名称。
"""


def ent(table="ods_trade_order_di", alias="a", target="dwb_trade_order_d"):
    return {"源表schema": "ods", "源表物理表名": table, "源表别名": alias,
            "目标表逻辑schema": "dws", "目标表物理名称": target,
            "关联&限定条件": "", "备注": ""}


def attr(tcol="channel_name", tcol_cn="渠道名称", ttype="VARCHAR(64)",
         alias="c", stable="dim_channel", scol="channel_name", group="default"):
    return {"分组": group, "源表别名": alias, "源表物理表名": stable, "源表字段名": scol,
            "映射规则": "直取", "映射表达式": "-", "目标字段名": tcol,
            "目标字段中文名": tcol_cn, "目标字段类型": ttype, "备注": ""}


def chg(seq, time_, page, key, column, ctype, before="", after=""):
    """变更记录行：page='entity'|'attr'；key='分组|名'；column='资产字段' 表示整行新增。"""
    page_cn = "实体级" if page == "entity" else "属性级"
    return {"变更序号": seq, "变更时间": time_, "变更页签": page_cn,
            "变更KEY": key, "变更列": column, "变更类型": ctype,
            "变更前": before, "变更后": after}


@pytest.fixture(scope="module")
def facts():
    demo = json.loads((FIXTURES / "baseline_v1_demo_full.json").read_text(encoding="utf-8"))
    ts, _ = build_ts_baseline(demo)
    from preprocess_opt import baseline_facts
    return baseline_facts(ts)


def make_pkg(pkg_dir: Path, entity_rows, attr_rows, change_rows, rs_md=RS_MD,
             mapping_name=PKG_XLSX, with_change_sheet=True):
    pkg_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws1 = wb.active; ws1.title = "实体级mapping"; ws1.append(ENTITY_COLS)
    for r in entity_rows:
        ws1.append([r.get(c, "") for c in ENTITY_COLS])
    ws2 = wb.create_sheet("属性级mapping"); ws2.append(ATTR_COLS)
    for r in attr_rows:
        ws2.append([r.get(c, "") for c in ATTR_COLS])
    if with_change_sheet:
        ws3 = wb.create_sheet("变更记录"); ws3.append(CHANGE_COLS)
        for r in change_rows:
            ws3.append([r.get(c, "") for c in CHANGE_COLS])
    wb.save(pkg_dir / mapping_name)
    (pkg_dir / PKG_RS).write_text(rs_md, encoding="utf-8")
    return pkg_dir


ATTR_EN = {"分组": "scene_group", "源表别名": "source_alias", "源表物理表名": "source_table",
           "源表字段名": "source_column", "映射规则": "mapping_rule", "映射表达式": "mapping_expression",
           "目标字段名": "target_column", "目标字段中文名": "target_column_cn",
           "目标字段类型": "target_type", "备注": "remark"}
ENT_EN = {"源表schema": "source_schema", "源表物理表名": "source_table", "源表别名": "source_alias",
          "目标表逻辑schema": "target_schema", "目标表物理名称": "target_table",
          "关联&限定条件": "join_condition", "备注": "remark"}


def mapping_of(entity_rows, attr_rows):
    """中文列名行 → extract_and_check 消费的英文键行（与 parse_mapping_sheets 归一后同构）。"""
    to_en = lambda rows, m: [{m.get(k, k): v for k, v in r.items()} for r in rows]
    return {"entity": to_en(entity_rows, ENT_EN), "attr": to_en(attr_rows, ATTR_EN)}


STD_CHANGES = [
    chg(1, "2026-03-10", "attr", "default|pay_type", "资产字段", "ADD"),          # 历史批次（不进本次）
    chg(2, "2026-10-14", "attr", "default|channel_name", "资产字段", "ADD"),      # 本次新增字段
]


def _to_entry(c: dict) -> dict:
    """chg 中文行 → read_change_sheet 归一后的 entry 同构形态（extract 单测直喂）——
    含标识符小写归一（与生产 read_change_sheet 同口径）。"""
    from preprocess_opt import normalize_yyyymm
    group, _, key_name = c["变更KEY"].partition("|")
    return {"seq": c["变更序号"], "time": normalize_yyyymm(c["变更时间"]),
            "page": "entity" if "实体" in c["变更页签"] else "attr",
            "group": group.strip().lower(), "key": key_name.strip().lower(),
            "column": c["变更列"],
            "ctype": c["变更类型"], "before": c["变更前"], "after": c["变更后"]}


class TestNormAsset:
    """资产名 I/F 镜像归一：资产铆定 I，mapping 写 I 视图、baseline 记 F 表是同一资产。"""

    def test_if_mirror_equal(self):
        assert norm_asset("dwb_trade_order_i") == norm_asset("dwb_trade_order_f")

    def test_schema_prefix_and_case(self):
        assert norm_asset("DWS.DWB_TRADE_ORDER_I") == norm_asset("dwb_trade_order")

    def test_bare_name(self):
        assert norm_asset("dwb_trade_order") == norm_asset("dwb_trade_order_i")

    def test_real_mismatch(self):
        assert norm_asset("dwb_cust_asset_i") != norm_asset("dwb_trade_order_f")


class TestChangeSheetParsing:
    """变更记录页签读取：八列规格/枚举校验/批次锚定/历史批次过滤。"""

    def test_read_and_batch(self, tmp_path):
        make_pkg(tmp_path, [ent(), ent(table="dim_channel", alias="c")],
                 [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"), attr()],
                 STD_CHANGES)
        entries = read_change_sheet(tmp_path / PKG_XLSX)
        assert len(entries) == 2
        assert pick_batch_version(entries) == "202610"
        batch = [e for e in entries if e["time"] == "202610"]
        assert [e["key"] for e in batch] == ["channel_name"]
        assert batch[0]["column"] == "资产字段" and batch[0]["ctype"] == "ADD"
        assert batch[0]["group"] == "default"

    def test_missing_sheet_rejected(self, tmp_path):
        make_pkg(tmp_path, [ent()], [attr()], STD_CHANGES, with_change_sheet=False)
        with pytest.raises(ValueError, match="变更记录"):
            read_change_sheet(tmp_path / PKG_XLSX)

    def test_bad_ctype_rejected(self, tmp_path):
        make_pkg(tmp_path, [ent()], [attr()],
                 [chg(1, "2026-10-14", "attr", "default|x", "资产字段", "新增")])
        with pytest.raises(ValueError, match="非法"):
            read_change_sheet(tmp_path / PKG_XLSX)

    def test_bad_time_rejected(self, tmp_path):
        make_pkg(tmp_path, [ent()], [attr()],
                 [chg(1, "不知道", "attr", "default|x", "资产字段", "ADD")])
        with pytest.raises(ValueError, match="归一失败"):
            read_change_sheet(tmp_path / PKG_XLSX)

    def test_column_name_case_insensitive(self, tmp_path):
        """列名大小写混合形态（变更Key）也能命中标准列（真实模板为「变更KEY」大写）。"""
        pkg = make_pkg(tmp_path, [ent(), ent(table="dim_channel", alias="c")], [attr()],
                       [chg(2, "2026-10-14", "attr", "default|channel_name", "资产字段", "ADD")])
        from openpyxl import load_workbook
        wb = load_workbook(pkg / PKG_XLSX)
        ws = wb["变更记录"]
        ws.cell(row=1, column=4).value = "变更Key"      # 标准名外的混合大小写
        wb.save(pkg / PKG_XLSX)
        entries = read_change_sheet(pkg / PKG_XLSX)
        assert entries[0]["key"] == "channel_name"

    def test_uppercase_key_normalized(self, tmp_path):
        """真实填写形态：变更 key 全大写 → read_change_sheet 归一为小写（标识符 foldcase）。"""
        make_pkg(tmp_path, [ent(), ent(table="dim_channel", alias="c")], [attr()],
                 [chg(2, "2026-10-14", "attr", "DEFAULT|CHANNEL_NAME", "资产字段", "ADD")])
        entries = read_change_sheet(tmp_path / PKG_XLSX)
        assert entries[0]["key"] == "channel_name" and entries[0]["group"] == "default"


class TestRenameSynth:
    """改名合成：严格证据链（DELETE.变更后 == ADD.key）才合成，独立行不误合。"""

    def test_paired_delete_add_synthesizes(self):
        batch = [_to_entry(chg(3, "2026-10-14", "attr", "default|order_amt", "目标字段名", "DELETE",
                               "order_amt", "order_amount")),
                 _to_entry(chg(4, "2026-10-14", "attr", "default|order_amount", "资产字段", "ADD"))]
        renames, rest = synth_renames(batch)
        assert renames == [{"kind": "rename", "group": "default",
                            "old": "order_amt", "new": "order_amount", "seq": "3+4",
                            "after_raw": "order_amount"}]
        assert rest == []

    def test_rename_with_uppercase_key_and_after(self):
        """真实填写形态：变更 key 大写、变更后人填值大写——标识符 foldcase 归一后仍配对。"""
        batch = [_to_entry(chg(3, "2026-10-14", "attr", "DEFAULT|ORDER_AMT", "目标字段名", "DELETE",
                               "ORDER_AMT", "ORDER_AMOUNT")),
                 _to_entry(chg(4, "2026-10-14", "attr", "DEFAULT|ORDER_AMOUNT", "资产字段", "ADD"))]
        renames, rest = synth_renames(batch)
        assert renames and renames[0]["old"] == "order_amt" and renames[0]["new"] == "order_amount"
        assert renames[0]["after_raw"] == "ORDER_AMOUNT"
        assert rest == []

    def test_unpaired_stays_independent(self):
        batch = [_to_entry(chg(1, "2026-10-14", "attr", "default|old_f", "目标字段名", "DELETE",
                               "old_f", "")),               # 变更后为空——无配对证据
                 _to_entry(chg(2, "2026-10-14", "attr", "default|new_f", "资产字段", "ADD"))]
        renames, rest = synth_renames(batch)
        assert renames == [] and len(rest) == 2

    def test_delete_after_points_elsewhere_no_synth(self):
        batch = [_to_entry(chg(1, "2026-10-14", "attr", "default|a", "目标字段名", "DELETE", "a", "zzz")),
                 _to_entry(chg(2, "2026-10-14", "attr", "default|b", "资产字段", "ADD"))]
        assert synth_renames(batch)[0] == []


class TestExtract:
    """提取与校验：add_field 候选/类型归类待扩展/conflict/dangling/漏标。"""

    def _std(self, extra_changes=(), extra_entity=(), extra_attr=()):
        entity = [ent(), ent(table="dim_channel", alias="c")] + list(extra_entity)
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"),
                 attr()] + list(extra_attr)
        mapping = mapping_of(entity, attrs)
        all_changes = STD_CHANGES + list(extra_changes)
        batch = [_to_entry(e) for e in all_changes if "10-14" in e["变更时间"]]
        return mapping, batch

    def test_extract_add_field(self, facts):
        mapping, batch = self._std()
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert [f["field"] for f in add_fields] == ["channel_name"]
        f = add_fields[0]
        assert f["source"]["table"] == "dim_channel" and f["cn"] == "渠道名称"
        assert f["new_source_table"] is True
        assert not [d for d in diags if d["level"] == "error"]

    def test_extract_add_field_uppercase_key(self, facts):
        """真实填写形态：变更 key 全大写（DEFAULT|CHANNEL_NAME）也能命中属性级小写行。"""
        entity = [ent(), ent(table="dim_channel", alias="c")]
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"), attr()]
        mapping = mapping_of(entity, attrs)
        batch = [_to_entry(chg(2, "2026-10-14", "attr", "DEFAULT|CHANNEL_NAME", "资产字段", "ADD"))]
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert [f["field"] for f in add_fields] == ["channel_name"], \
            "大写 key 归一后应命中属性级小写行"
        assert not [d for d in diags if d["level"] == "error"]

    def test_modify_delete_classified_unsupported(self, facts):
        """非目标字段类型列的 MODIFY/DELETE/列级 ADD 仍待扩展；实体级 ADD（场景 A）
        已正式消费——当前态无行时报 add_source_row_missing error（不进 unsupported）。"""
        extra = [chg(3, "2026-10-14", "attr", "default|order_id", "映射规则", "MODIFY"),
                 chg(4, "2026-10-14", "attr", "default|gone_f", "资产字段", "DELETE"),
                 chg(5, "2026-10-14", "entity", "default|dim_shop", "资产字段", "ADD"),
                 chg(6, "2026-10-14", "attr", "default|order_id", "备注", "ADD")]
        mapping, batch = self._std(extra)
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        kinds = {(u["change_type"], u["name"]) for u in unsupported}
        assert ("modify", "order_id") in kinds
        assert ("delete", "gone_f") in kinds
        assert ("add_column", "order_id.备注") in kinds
        assert ("add_source", "dim_shop") not in kinds, "实体级 ADD 已正式消费"
        assert any(d["code"] == "add_source_row_missing" and "dim_shop" in d["message"]
                   for d in diags if d["level"] == "error")

    def test_rename_reported_unsupported(self, facts):
        extra = [chg(3, "2026-10-14", "attr", "default|order_amt", "目标字段名", "DELETE",
                     "order_amt", "order_amount"),
                 chg(4, "2026-10-14", "attr", "default|order_amount", "资产字段", "ADD")]
        mapping, batch = self._std(extra)
        renames, rest = synth_renames(batch)
        assert renames, "先删后加对应合成 rename"
        _, _ns, _mods, _mlogics, unsupported, _ = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert any(u["change_type"] == "rename" and "order_amt→order_amount" in u["name"]
                   for u in unsupported)

    def test_conflict_and_dangling_and_row_missing(self, facts):
        mapping, batch = self._std(extra_changes=[
            chg(7, "2026-10-14", "attr", "default|order_id", "资产字段", "ADD"),   # baseline 已有
            chg(8, "2026-10-14", "attr", "default|ghost_f", "资产字段", "ADD"),    # 当前态无行
        ])
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        codes = {d["code"] for d in diags if d["level"] == "error"}
        assert "add_field_conflict" in codes
        assert "add_field_row_missing" in codes
        assert [f["field"] for f in add_fields] == ["channel_name"]

    def test_alias_dangling(self, facts):
        mapping, batch = self._std(extra_changes=[
            chg(9, "2026-10-14", "attr", "default|dangle_f", "资产字段", "ADD"),
        ], extra_attr=[attr(tcol="dangle_f", alias="nope", stable="dim_x", scol="x")])
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert any(d["code"] == "source_alias_dangling" for d in diags if d["level"] == "error")

    def test_unmarked_new_field_warns(self, facts):
        # 属性级当前态有 extra_f 不在 baseline 且不在变更记录 → 漏标 warn（台账漂移防线）
        mapping, batch = self._std(extra_attr=[
            attr(tcol="extra_f", alias="a", stable="ods_trade_order_di", scol="extra_f")])
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert any(d["code"] == "unmarked_new_field" and "extra_f" in d["message"]
                   for d in diags if d["level"] == "warn")

    def test_rs_material_reconciliation(self, facts):
        mapping, batch = self._std()
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610",
                                        "新增 other_f 字段")   # RS 段没提 channel_name
        assert any(d["code"] == "rs_field_not_mentioned" for d in diags if d["level"] == "warn")


class TestRsParsing:
    """RS 材料角色（v3：版本锚定主源已移至 mapping 变更记录页签）。"""

    def test_version_section_anchor(self):
        rs = "前文\n202610版本 需求正文A\n202603版本 需求正文B\n"
        assert "需求正文A" in extract_version_section(rs, "202610")
        assert "202603版本" not in extract_version_section(rs, "202610")

    def test_parse_change_log_material(self):
        rs = "## 3.3 变更记录\n\n| 日期 | 版本 | 修改人 | 修改内容 |\n|---|---|---|---|\n| 2026-10-14 | v3 | x | 优化 |\n"
        assert parse_change_log(rs)[0]["date"] == "2026-10-14"


def _opt_ready(root, baseline_ts):
    arc = root / "archive"
    arc.mkdir(parents=True, exist_ok=True)
    (arc / "ts.json").write_text(json.dumps(baseline_ts, ensure_ascii=False), encoding="utf-8")
    return arc


@pytest.fixture()
def demo_baseline():
    demo = json.loads((FIXTURES / "baseline_v1_demo_full.json").read_text(encoding="utf-8"))
    ts, _ = build_ts_baseline(demo)
    return ts


class TestMainEndToEnd:
    def test_direct_paths_flow(self, tmp_path, demo_baseline):
        """e2e 含场景 A（新来源表声明）：变更记录带实体级 ADD（dim_channel）+
        属性级新字段（来自该表）——联动两向对齐无 warn。"""
        arc = _opt_ready(tmp_path, demo_baseline)
        changes = STD_CHANGES + [chg(3, "2026-10-14", "entity", "default|c", "资产字段", "ADD")]
        pkg = make_pkg(tmp_path / "pkg",
                       [ent(), ent(table="dim_channel", alias="c")],
                       [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"),
                        attr()],
                       changes)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc in (0, 1)   # 1=G2 台账缺页披露（最小夹具不声明存量字段——预期）
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        assert cr["version"] == "202610"
        assert [f["field"] for f in cr["fields"]] == ["channel_name"]
        assert [ns["table"] for ns in cr["new_sources"]] == ["dim_channel"]
        assert len(cr["change_entries"]) == 2          # 本批次（实体级 ADD + 属性级 ADD）
        assert "mapping" in cr["source_files"]

    def test_explicit_version_override_picks_history_batch(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        pkg = make_pkg(tmp_path / "pkg", [ent()],
                       [attr(tcol="pay_type", alias="a", stable="ods_trade_order_di", scol="pay_type"),
                        attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id")],
                       STD_CHANGES)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path),
                   "--version", "202603"])
        assert rc in (0, 1)
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        assert cr["version"] == "202603"
        assert [f["field"] for f in cr["fields"]] == ["pay_type"]

    def test_blocked_exit_2(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        pkg = make_pkg(tmp_path / "pkg",
                       [ent(target="dwb_other_asset_d")],     # 资产不一致 → 阻断
                       [attr()], STD_CHANGES)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc == 2

    def test_missing_input_file_exit_2(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        rc = main(["--mapping", str(tmp_path / "nope.xlsx"), "--rs", str(tmp_path / "nope.md"),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc == 2


class TestBaselineViewBackfill:
    def test_view_rendered_when_missing(self):
        from preprocess_opt import render_baseline_view_from_ts
        ts = {"meta": {"target": {"f_table": {"schema": "dws", "table": "t_f"}}},
              "rules": {"R0001": {"exec_sequence": 1, "load_mode": "no_delete",
                                  "target_table": "dws.t_f",
                                  "source_tables": [{"schema": "ods", "table": "s"}]}},
              "tables": {"t_f": {"fields": [{"name": "a"}, {"name": "b"}]}}}
        view = render_baseline_view_from_ts(ts)
        assert "R0001" in view and "t_f" in view and "2 字段" in view


class TestBaselineViewDecisionsLanguage:
    """decisions 语言地图（2026-10-23 批1）：判断位有则显、无则标空——adopt/逆向档差异
    由数据存在性自然呈现（渲染器不分叉）；territory（逐字段血缘）不进 designer 输入面。"""

    def _adopt_ts(self):
        return {"meta": {"target": {"f_table": {"schema": "dws", "table": "t_f"}},
                         "grain": {"input": "一行一单", "output": "一行一单"},
                         "business_key": ["order_id"]},
                "rules": {"R0001": {"exec_sequence": 1, "load_mode": "merge_into",
                                    "step_type": "merge", "target_role": "target",
                                    "target_table": "dws.t_f",
                                    "source_tables": [{"schema": "ods", "table": "s", "alias": "a"}],
                                    "joins": [{"alias": "c", "type": "LEFT", "condition": "a.id=c.id"}],
                                    "join_safety": [{"alias": "c", "join_key_unique": True,
                                                     "strategy": "直接关联", "reason": "维表主键关联"}]}},
                "tables": {"t_f": {"fields": [{"target_field": "x"}]}}}

    def test_adopt_judgments_rendered(self):
        from preprocess_opt import render_baseline_view_from_ts
        view = render_baseline_view_from_ts(self._adopt_ts())
        assert "一行一单 → 一行一单" in view              # grain dict 形态归一显示
        assert "business_key: order_id" in view
        assert "唯一=True·直接关联（维表主键关联）" in view  # join_safety 意图随档可见
        assert "▢未声明" not in view                      # 真档判断位全显无空位（header 说明字样除外）
        assert "字段血缘摘要" not in view                  # territory 不渲染

    def test_reverse_gaps_marked(self):
        from preprocess_opt import render_baseline_view_from_ts
        ts = {"meta": {"target": {"f_table": {"schema": "dws", "table": "t_f"}}},
              "rules": {"R0001": {"target_table": "t_f",
                                  "source_tables": [{"schema": "ods", "table": "s", "alias": "a"}]}},
              "tables": {}}
        view = render_baseline_view_from_ts(ts)
        assert "▢未声明（语义空位）" in view               # 粒度/主键空位标注
        assert "JOIN: （无声明）" in view

    def test_merge_on_evidence_when_bk_empty(self):
        """批0-③：business_key 空位时 MERGE ON 键旁证随行（强旁证——upsert 行身份）；
        判断位不填防消费者误信。有 bk 则不出旁证行。"""
        from preprocess_opt import render_baseline_view_from_ts
        ts = {"meta": {"target": {"f_table": {"schema": "dws", "table": "t_f"}}},
              "rules": {"R0001": {"target_table": "dws.t_f", "load_mode": "merge_into",
                                  "write_condition": "T.order_id=T1.order_id"}},
              "tables": {}}
        view = render_baseline_view_from_ts(ts)
        assert "business_key: ▢未声明（语义空位）" in view
        assert "旁证（强·MERGE ON" in view and "R0001·T.order_id=T1.order_id" in view
        ts["meta"]["business_key"] = ["order_id"]
        view2 = render_baseline_view_from_ts(ts)
        assert "旁证" not in view2                          # 有声明不出旁证


class TestBuildAliasMap:
    """别名对照（落笔权威=SQL 实际别名）：确定性查表归脚本——designer 禁读 SQL 原文。"""

    def test_match_mismatch_missing(self):
        from preprocess_opt import build_alias_map
        ts = {"rules": {
            "R0001": {"source_tables": [{"table": "ods_a", "alias": "a"},
                                        {"table": "ods_b", "alias": "b"}]},
            "R0002": {"source_tables": [{"table": "ods_c", "alias": "c"}]}}}
        sql = {"R0001": "SELECT a.x FROM ods.ods_a a JOIN ods.ods_b ob ON a.id=ob.id"}
        rows = build_alias_map(sql, ts)
        by = {(r["rule"], r["table"]): r for r in rows}
        assert by[("R0001", "ods_a")]["sql_alias"] == "a"
        assert by[("R0001", "ods_a")]["note"] == ""       # 一致无标注
        assert by[("R0001", "ods_b")]["sql_alias"] == "ob"
        assert "以 SQL 为准" in by[("R0001", "ods_b")]["note"]
        assert "SQL 缺失" in by[("R0002", "ods_c")]["note"]

    def test_cte_refs_not_noise(self):
        from preprocess_opt import build_alias_map
        ts = {"rules": {"R0001": {"source_tables": [{"table": "ods_a", "alias": "a"}]}}}
        sql = {"R0001": "WITH t AS (SELECT 1 FROM ods.ods_a a) SELECT * FROM t x"}
        rows = build_alias_map(sql, ts)
        assert [r["table"] for r in rows] == ["ods_a"]     # CTE 名不进对照


class TestBuildPlacementHints:
    """落位预判（确定性事实）：直挂/穿中间表/新来源三态——取舍判断仍归 designer。"""

    def _ts(self):
        return {"meta": {"target": {"f_table": {"table": "t_f"}}},
                "rules": {
                    "R0001": {"target_table": "dws.tmp_x",
                              "source_tables": [{"table": "ods_a", "alias": "a"}]},
                    "R0002": {"target_table": "dws.t_f", "reads": ["dws.tmp_x"],
                              "source_tables": [{"table": "ods_b", "alias": "b"}]}}}

    def test_direct_via_new(self):
        from preprocess_opt import build_placement_hints
        fields = [
            {"field": "f1", "source": {"table": "ods_b"}},             # 直挂 R0002
            {"field": "f2", "source": {"table": "ods_a"}},             # 穿中间表 R0001→R0002
            {"field": "f3", "source": {"table": "dim_new"}, "new_source_table": True},
        ]
        hints = {h["field"]: h for h in build_placement_hints(self._ts(), fields)}
        assert hints["f1"]["direct_rules"] == ["R0002"]
        assert hints["f1"]["via_intermediate"] == {}
        assert hints["f2"]["direct_rules"] == []
        assert hints["f2"]["via_intermediate"] == {"R0001": ["R0002"]}
        assert hints["f3"]["direct_rules"] == []
        assert hints["f3"]["via_intermediate"] == {}
        assert hints["f3"]["new_source"] is True
        assert "新 JOIN" in hints["f3"]["note"]


class TestInputSurfaceE2E:
    """e2e：落位预判进 change_request、别名对照进 baseline_view（批1 输入面重构）。"""

    def test_placement_hints_and_alias_map_e2e(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        # 档案 SQL：R0001 实际别名故意与 ts 声明不同名（ord≠a）——对照表应抓到
        (arc / "etl").mkdir(exist_ok=True)
        (arc / "etl" / "R0001.sql").write_text(
            "SELECT ord.order_id FROM ods.ods_trade_order_di ord;", encoding="utf-8")
        changes = STD_CHANGES + [
            chg(3, "2026-10-14", "entity", "default|c", "资产字段", "ADD"),
            chg(4, "2026-10-14", "attr", "default|pay_type", "资产字段", "ADD")]
        rs_md = "# 某资产需求规格\n\n## 4. 字段需求\n\n202610版本新增需求：新增 channel_name（渠道名称），来源渠道维表；新增 pay_type（支付方式）。\n"
        pkg = make_pkg(tmp_path / "pkg",
                       [ent(), ent(table="dim_channel", alias="c")],
                       [attr(tcol="pay_type", alias="a", stable="ods_trade_order_di", scol="pay_type"),
                        attr()],
                       changes, rs_md=rs_md)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc in (0, 1)   # 1=G2 披露（最小夹具不声明存量字段——预期）
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        hints = {h["field"]: h for h in cr["placement_hints"]}
        # pay_type 源表 ods_trade_order_di：只可见于中间规则 R0001（产 tmp）——穿中间表路径
        assert hints["pay_type"]["direct_rules"] == []
        assert hints["pay_type"]["via_intermediate"] == {"R0001": ["R0002"]}
        # channel_name 新来源（dim_channel）——新 JOIN 提示
        assert hints["channel_name"]["new_source"] is True
        assert "新 JOIN" in hints["channel_name"]["note"]
        # 别名对照进 view：SQL 实际别名 ord ≠ ts 声明 a → 不一致标注（落笔以 SQL 为准）
        view = (tmp_path / "build" / "_internal" / "baseline_view.md").read_text(encoding="utf-8")
        assert "## 别名对照" in view
        assert "ord" in view and "以 SQL 为准" in view
        # 拓扑段：穿中间表路径（reads）在图上可见
        assert "### R0001 → dws.tmp_trade_order" in view
        assert "读中间表: dws.tmp_trade_order" in view


class TestEvalWorksheet:
    """评估工作单+切片上下文（批2 断言对齐）：新 JOIN 断言须有据——探针引擎零改动，
    切片=确定性适配器（最小 rs_input 形态只含新来源表）。无新来源不产。"""

    def test_eval_context_and_worksheet_e2e(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        ent_c = ent(table="dim_channel", alias="c")
        ent_c["源表schema"] = "dws"
        ent_c["关联&限定条件"] = "t.order_id = c.order_id"   # 实体级声明的关联条件→预填行
        changes = STD_CHANGES + [chg(3, "2026-10-14", "entity", "default|c", "资产字段", "ADD")]
        pkg = make_pkg(tmp_path / "pkg", [ent(), ent_c], [attr()], changes)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc in (0, 1)   # 1=G2 披露（最小夹具——预期）
        ctx_path = tmp_path / "build" / "_internal" / "eval_context.json"
        assert ctx_path.exists()
        ctx = json.loads(ctx_path.read_text(encoding="utf-8"))
        assert ctx["meta"]["target"]["f_table"]["schema"] == "dws"   # 探针选源锚点
        assert ctx["source_tables"] == [{"source_alias": "c", "source_schema": "dws",
                                         "source_table": "dim_channel",
                                         "join_condition": "t.order_id = c.order_id"}]
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        # 工作单预填行：结构化条件预解析（c 侧键=order_id），designer 核一眼跑实测
        assert "c|order_id|" in cr["eval_worksheet"]

    def test_no_new_source_no_worksheet(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        changes = [chg(1, "2026-10-14", "attr", "default|pay_type", "资产字段", "ADD")]
        rs_md = "# RS\n\n## 4. 字段需求\n\n202610版本新增 pay_type。\n"
        pkg = make_pkg(tmp_path / "pkg", [ent()],
                       [attr(tcol="pay_type", alias="a", stable="ods_trade_order_di",
                             scol="pay_type")], changes, rs_md=rs_md)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc in (0, 1)   # 1=G2 披露（最小夹具——预期）
        assert not (tmp_path / "build" / "_internal" / "eval_context.json").exists()
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        assert cr["eval_worksheet"] == ""


class TestHealthReport:
    """批1 体检报告（跨脚本短路修复+G1/G2 披露）：阻断仍落 change_request（_blocked_
    errors 随件内联，precheck 降级并入一份报告）+ 新检查默认披露级。"""

    def test_blocked_still_writes_cr(self, tmp_path, demo_baseline):
        """资产不一致（阻断）→ change_request 仍落盘带 _blocked_errors——人一次收全。"""
        arc = _opt_ready(tmp_path, demo_baseline)
        pkg = make_pkg(tmp_path / "pkg", [ent(target="dwb_other_asset_d")], [attr()],
                       STD_CHANGES)
        rc = main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
                   "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert rc == 2
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        assert cr["_blocked"] is True
        assert any(e["code"] == "asset_table_mismatch" for e in cr["_blocked_errors"])

    def test_g1_delete_row_still_present_disclosed(self, tmp_path, demo_baseline):
        """G1：变更记录标 DELETE 但当前态该行仍在 → 披露 warn（台账时间线存疑）。"""
        arc = _opt_ready(tmp_path, demo_baseline)
        changes = [chg(1, "2026-10-14", "attr", "default|order_id", "资产字段", "DELETE")]
        pkg = make_pkg(tmp_path / "pkg", [ent()],
                       [attr(tcol="order_id", alias="a", stable="ods_trade_order_di",
                             scol="order_id")], changes)
        main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
             "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        assert any(w["code"] == "delete_row_still_present" for w in cr["_warns"])

    def test_g2_ghost_fields_disclosed(self, tmp_path, demo_baseline):
        """G2：baseline 有、台账无、无 DELETE 行 → 台账缺页披露（漏标扫描的对称面）。"""
        arc = _opt_ready(tmp_path, demo_baseline)
        pkg = make_pkg(tmp_path / "pkg", [ent()], [attr()], STD_CHANGES)
        main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
             "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        cr = json.loads((tmp_path / "build" / "_internal" / "change_request.json").read_text())
        ghost = next(w for w in cr["_warns"] if w["code"] == "ledger_ghost_fields")
        assert "order_id" in ghost["message"]     # demo baseline 字段未声明

    def test_g3_crosscheck_in_view(self):
        """G3：DDL 列↔ETL 投影列差渲染进 view 披露段（两路径共用）。"""
        from preprocess_opt import render_baseline_view_from_ts
        ts = {"meta": {"target": {"f_table": {"schema": "dws", "table": "t_f"}}},
              "rules": {"R0001": {"target_table": "dws.t_f",
                                  "fields": {"direct": ["a.x AS x"]}}},
              "tables": {"t_f": {"fields": [{"target_field": "x"}, {"target_field": "y"}]}}}
        view = render_baseline_view_from_ts(ts)
        assert "制品交叉披露" in view and "无规则产出" in view and "y" in view


class TestPrecheckHealthReport:
    """precheck_opt 聚合报告：上游阻断并入+三层分组+diff 增量+残次产物 exit 2。"""

    def _cr(self, tmp_path, blocked=None, warns=None):
        out = tmp_path / "build" / "_internal"
        out.mkdir(parents=True, exist_ok=True)
        cr = {"fields": [], "modify_type_fields": [], "_blocked": bool(blocked),
              "_blocked_errors": blocked or [], "_warns": warns or []}
        (out / "change_request.json").write_text(json.dumps(cr, ensure_ascii=False),
                                                 encoding="utf-8")
        return out / "change_request.json"

    def test_blocked_cr_aggregated_and_exit2(self, tmp_path, demo_baseline):
        from precheck_opt import main as pc_main
        arc = _opt_ready(tmp_path, demo_baseline)
        cr_path = self._cr(tmp_path, blocked=[{"code": "asset_mismatch", "message": "串包"}])
        rc = pc_main(["--change-request", str(cr_path),
                      "--ts-baseline", str(arc / "ts.json"), "--outdir", str(cr_path.parent)])
        assert rc == 2                       # 残次 change_request 不进设计
        report = (cr_path.parent / "input_health_report.md").read_text(encoding="utf-8")
        assert "必须修" in report and "asset_mismatch" in report and "串包" in report

    def test_diff_increment_marks_new(self, tmp_path):
        from precheck_opt import write_health_report
        from risk_checks import PrecheckResult
        out = tmp_path / "build" / "_internal"
        out.mkdir(parents=True, exist_ok=True)
        cr_path = out / "change_request.json"
        # 上轮报告（ddlc/last_input_health.md——preprocess 清场前抢救位）
        (tmp_path / "last_input_health.md").write_text(
            "# 输入体检报告\n\n▶ 带病可继续\n\n- [ledger_ghost_fields] baseline 有 2 个字段不在\n",
            encoding="utf-8")
        r = PrecheckResult()
        r.add_warn("[ledger_ghost_fields] baseline 有 2 个字段不在")
        r.add_warn("[新检查] 本轮新冒的披露")
        write_health_report(out, [], [], r, cr_path)
        text = (out / "input_health_report.md").read_text(encoding="utf-8")
        assert "▲新增 1" in text and "▲ [新检查] 本轮新冒的披露" in text
        assert "✓消失 0" in text


class TestWorkspaceReady:
    def test_build_cleared_and_tmp_created(self, tmp_path, demo_baseline):
        arc = _opt_ready(tmp_path, demo_baseline)
        stale = tmp_path / "build" / "stale.txt"
        stale.parent.mkdir(parents=True, exist_ok=True)
        stale.write_text("旧现场残留", encoding="utf-8")
        pkg = make_pkg(tmp_path / "pkg", [ent(), ent(table="dim_channel", alias="c")],
                       [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"),
                        attr()], STD_CHANGES)
        main(["--mapping", str(pkg / PKG_XLSX), "--rs", str(pkg / PKG_RS),
              "--ts-baseline", str(arc / "ts.json"), "--opt-root", str(tmp_path)])
        assert not stale.exists()                       # 开工清场
        assert (tmp_path / "archive_tmp" / "ts.json").exists()   # 临时档案=档案副本起步


class TestFieldTypeDrift:
    """存量字段类型对账（同名异容，2026-10-21 三方一致性补位）+ unsupported 存在性事实。"""

    def test_drift_warns_on_type_change(self, facts):
        """mapping 改存量字段类型（VARCHAR(64)→128）无变更记录 → warn 披露两侧类型。"""
        from preprocess_opt import field_type_drift
        rows = [{"target_column": "order_id", "target_type": "VARCHAR(128)"}]
        diags = field_type_drift(rows, facts["field_types"], set())
        assert [d["code"] for d in diags] == ["field_type_drift"]
        assert "VARCHAR(128)" in diags[0]["message"] and "VARCHAR(64)" in diags[0]["message"]
        assert diags[0]["level"] == "warn"

    def test_normalized_form_not_warned(self, facts):
        """书写形态差异（NUMBER(18,2) vs 基线 DECIMAL(18,2)；小写 varchar）归一后不报。"""
        from preprocess_opt import field_type_drift
        rows = [{"target_column": "total_amount", "target_type": "NUMBER(18,2)"},
                {"target_column": "order_id", "target_type": "varchar(64)"}]
        assert field_type_drift(rows, facts["field_types"], set()) == []

    def test_family_change_warned(self, facts):
        from preprocess_opt import field_type_drift
        rows = [{"target_column": "order_id", "target_type": "numeric(20,0)"}]
        diags = field_type_drift(rows, facts["field_types"], set())
        assert [d["code"] for d in diags] == ["field_type_drift"]

    def test_skip_and_empty_sides(self, facts):
        """skip（本批新增/改名）跳过；任一侧类型空=宁放过；不在基线=漏标扫描的领地。"""
        from preprocess_opt import field_type_drift
        rows = [{"target_column": "order_id", "target_type": "VARCHAR(128)"},   # 命中 skip
                {"target_column": "cust_id", "target_type": ""},                # 声明侧空
                {"target_column": "channel_name", "target_type": "INT"}]       # 不在基线
        assert field_type_drift(rows, facts["field_types"], {"order_id"}) == []

    def test_extract_chain_emits_drift(self, facts):
        """真实链路：extract_and_check 对存量字段类型漂移产 warn（本批 ADD 不受影响）。"""
        extra_attr = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di",
                           scol="order_id", ttype="VARCHAR(200)")]
        mapping, batch = self._std(extra_attr=extra_attr) if hasattr(self, "_std") else (None, None)
        std = TestExtract()
        mapping, batch = std._std(extra_attr=extra_attr)
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        drift = [d for d in diags if d["code"] == "field_type_drift"]
        assert drift and "order_id" in drift[0]["message"]
        assert [f["field"] for f in add_fields] == ["channel_name"]  # 本批 ADD 照常

    def test_unsupported_existence_note(self, facts):
        """unsupported 报告的目标存在性事实：笔误/台账滞后/两侧都有（真待支持）三分。"""
        extra = [chg(3, "2026-10-14", "attr", "default|typo_f", "映射规则", "MODIFY"),      # 两侧都无（映射表达式列已被第三刀消费）
                 chg(4, "2026-10-14", "attr", "default|total_amount", "映射规则", "MODIFY"),  # 基线有当前态无
                 chg(5, "2026-10-14", "attr", "default|order_id", "备注", "ADD")]            # 两侧都有
        std = TestExtract()
        mapping, batch = std._std(extra_changes=extra)
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, unsupported, _ = extract_and_check(rest, renames, mapping, facts, "202610", "")
        notes = {u["name"]: u.get("note", "") for u in unsupported}
        assert "疑似笔误" in notes["typo_f"]
        assert "台账滞后" in notes["total_amount"]
        assert notes.get("order_id", "") == "" or "order_id.备注" not in notes
        # 列级 ADD 的 note 键名是 "order_id.备注"（key 是字段名——两侧都有 → 空 note）
        col_add = next(u for u in unsupported if u["change_type"] == "add_column")
        assert col_add.get("note", "") == ""

    def test_entity_unsupported_note(self, facts):
        """实体级 unsupported（仍待扩展的类型，如 DELETE）：key 未命中实体级当前态
        （别名/分组）→ 事实附加。"""
        extra = [chg(6, "2026-10-14", "entity", "default|ghost_src", "资产字段", "DELETE")]
        std = TestExtract()
        mapping, batch = std._std(extra_changes=extra)
        renames, rest = synth_renames(batch)
        _, _ns, _mods, _mlogics, unsupported, _ = extract_and_check(rest, renames, mapping, facts, "202610", "")
        ent = next(u for u in unsupported if u["level"] == "entity")
        assert "实体级当前态" in ent.get("note", "")


class TestMultiAliasAddField:
    """多别名形态（2026-10-22 内网反馈）：新增加工字段的源表别名列填多别名分号隔开
    （提示性登记不保证完整）——拆分号任一命中实体级即过；全不命中 → dangling error
    且**漏标扫描不误报**（声明集与通过集分离：变更记录明确写了 ADD 的字段不会被说
    "不在本次变更记录中"）。"""

    def _mapping_with(self, alias, scol="id;channel_name"):
        """channel_name 行用指定别名/字段配对形态（2026-10-22 配对自洽定调：
        别名与字段一一对应——别名 2 个则字段 2 个）。"""
        entity = [ent(), ent(table="dim_channel", alias="c")]
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"),
                 attr(tcol="channel_name", alias=alias, stable="dim_channel", scol=scol)]
        mapping = mapping_of(entity, attrs)
        batch = [_to_entry(c) for c in STD_CHANGES if "10-14" in c["变更时间"]]
        return mapping, batch

    def test_multi_alias_first_hit_used(self, facts):
        """配对 'a;b' × 'id;channel_name'（首对 a 命中实体级）→ 无 dangling；
        source.alias=首个命中 a。"""
        mapping, batch = self._mapping_with("a;b")
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, unsupported, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert not [d for d in diags if d["code"] == "source_alias_dangling"], diags
        assert add_fields and add_fields[0]["source"]["alias"] == "a"

    def test_qualified_form_pair_used(self, facts):
        """字段格自带限定符 'a.id;c.channel_name'（别名格空——限定符优先）→
        首对 a 命中，source.field 取配对列名（不再整格）。"""
        mapping, batch = self._mapping_with("", scol="a.id;c.channel_name")
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert not [d for d in diags if d["code"] == "source_alias_dangling"], diags
        assert add_fields[0]["source"]["alias"] == "a"
        assert add_fields[0]["source"]["field"] == "id"

    def test_full_width_semicolon_also_splits(self, facts):
        """全角分号配对 'a；c'（a 命中实体级）→ 命中即过。"""
        mapping, batch = self._mapping_with("a；c")
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        assert not [d for d in diags if d["code"] == "source_alias_dangling"], diags
        assert add_fields[0]["source"]["alias"] == "a"

    def test_dangling_not_misreported_as_unmarked(self, facts):
        """连带问题回归：配对全不命中 → dangling error 照报，但该字段是变更记录
        明确声明的 ADD——漏标扫描不得再说它'不在本次变更记录中'。"""
        mapping, batch = self._mapping_with("xx;yy")
        renames, rest = synth_renames(batch)
        add_fields, _ns, _mods, _mlogics, _, diags = extract_and_check(rest, renames, mapping, facts, "202610", "")
        dangling = [d for d in diags if d["code"] == "source_alias_dangling"]
        unmarked = [d for d in diags if d["code"] == "unmarked_new_field"
                    and "channel_name" in d["message"]]
        assert dangling, "全不命中应报 dangling"
        assert not unmarked, f"变更记录已声明 ADD 的字段不得被漏标误报: {unmarked}"


class TestMappingSheetsConverged:
    """opt 全量 mapping 解析收敛（2026-10-22）：read_full_mapping 退役——
    内网坑回归锁：大写列头「源表Schema」在 opt 链路同样命中（此前 source.schema
    恒空一路漏到 change_request）。"""

    def _bigcase_pkg(self, tmp_path):
        """列头大写驼峰形态的真实风 mapping（用户现场形态）。"""
        entity_cols = ["源表Schema", "源表物理表名", "源表中文名", "源表别名",
                       "目标表逻辑Schema", "目标表中文名", "目标表物理名称", "关联&限定条件", "备注"]
        attr_cols = ["源Schema", "源表物理表名", "源表别名", "源表字段名", "源表字段中文名",
                     "源表字段类型", "映射规则", "映射表达式", "目标字段名", "目标字段中文名",
                     "目标字段类型", "备注"]
        pkg = tmp_path / "pkg"
        pkg.mkdir(parents=True)
        wb = Workbook()
        ws1 = wb.active; ws1.title = "实体级mapping"; ws1.append(entity_cols)
        ws1.append(["ods", "ods_trade_order_di", "订单", "a", "dws", "宽表", "dwb_trade_order_d", "", ""])
        ws2 = wb.create_sheet("属性级mapping"); ws2.append(attr_cols)
        ws2.append(["ods", "ods_trade_order_di", "a", "order_id", "订单号", "bigint",
                    "直取", "-", "order_id", "订单号", "bigint", ""])
        ws3 = wb.create_sheet("变更记录"); ws3.append(CHANGE_COLS)
        ws3.append([1, "2026-10-14", "属性级", "default|order_id", "备注", "MODIFY", "", ""])
        wb.save(pkg / PKG_XLSX)
        (pkg / PKG_RS).write_text(RS_MD, encoding="utf-8")
        return pkg

    def test_uppercase_columns_flow_through_opt(self, tmp_path, monkeypatch):
        """大写列头 mapping 走 opt 预处理全链：解析/静态检查/变更匹配均正常
        （不再 source_schema 恒空/缺列误炸）。"""
        import preprocess_opt as popt
        pkg = self._bigcase_pkg(tmp_path)
        entries = popt.read_change_sheet(pkg / PKG_XLSX)
        assert entries and entries[0]["key"] == "order_id"
        mapping = popt.parse_mapping_sheets(pkg / PKG_XLSX)
        assert mapping["entity"][0]["source_schema"] == "ods", "大写 Schema 列头必须命中"
        assert mapping["attr"][0]["source_schema"] == "ods"
        diags = popt.check_mapping_static(mapping["attr"], mapping["entity"])
        assert not [d for d in diags if d["level"] == "error"], diags

    def test_static_check_shared_semantics_in_opt(self, tmp_path):
        """opt 接静态检查=同一实现同一语义：直接复制字段别名不在实体级 → error 阻断。"""
        import preprocess_opt as popt
        from preprocess import parse_mapping_sheets, check_mapping_static
        pkg = self._bigcase_pkg(tmp_path)
        mapping = parse_mapping_sheets(pkg / PKG_XLSX)
        mapping["attr"][0]["source_alias"] = "ghost"  # 不在实体级
        diags = check_mapping_static(mapping["attr"], mapping["entity"])
        assert any(d["level"] == "error" and "ghost" in d["message"] for d in diags)


def facts_fixture():
    """TestScenario* 直调 extract_and_check 用的 facts（与 module fixture 同款构造）。"""
    demo = json.loads((FIXTURES / "baseline_v1_demo_full.json").read_text(encoding="utf-8"))
    ts, _ = build_ts_baseline(demo)
    from preprocess_opt import baseline_facts
    return baseline_facts(ts)


class TestScenarioNewSource:
    """场景 A（2026-10-22 第二刀·内网第一案例综合场景）：新增来源表声明
    （实体级整行 ADD）——正式消费进 new_sources + 与新字段引用两向联动。"""

    def _run(self, extra_changes=(), extra_entity=(), extra_attr=(), facts_=None):
        entity = [ent(), ent(table="dim_channel", alias="c")] + list(extra_entity)
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"),
                 attr()] + list(extra_attr)
        mapping = mapping_of(entity, attrs)
        all_changes = STD_CHANGES + list(extra_changes)
        batch = [_to_entry(c) for c in all_changes if "10-14" in c["变更时间"]]
        renames, rest = synth_renames(batch)
        return extract_and_check(rest, renames, mapping, facts_ or facts_fixture(), "202610", "")

    def test_declared_source_collected(self):
        add, new_sources, mods, mlogics, uns, diags = self._run(
            extra_changes=[chg(3, "2026-10-14", "entity", "default|c", "资产字段", "ADD")])
        assert [ns["table"] for ns in new_sources] == ["dim_channel"]
        assert not [d for d in diags if d["level"] == "error"], diags
        assert not [d for d in diags if d["code"] in ("new_source_undeclared", "new_source_unreferenced")]

    def test_declared_but_in_baseline_warns(self):
        """声明的'新'表其实已在 baseline 来源 → warn（重复申报）。"""
        add, new_sources, mods, mlogics, uns, diags = self._run(
            extra_changes=[chg(3, "2026-10-14", "entity", "default|a", "资产字段", "ADD")])
        assert any(d["code"] == "add_source_already_in_baseline" for d in diags)

    def test_reference_without_declaration_warns(self):
        """新字段来自新表但无实体级 ADD 声明 → warn（隐式放行，建议补申报）。"""
        add, new_sources, mods, mlogics, uns, diags = self._run()
        assert any(d["code"] == "new_source_undeclared" and "'c'" in d["message"] for d in diags)

    def test_declaration_without_reference_warns(self):
        """声明了新表（u=dim_unused）但无字段引用它 → warn（无用声明）。"""
        entity_extra = [ent(table="dim_unused", alias="u")]
        add, new_sources, mods, mlogics, uns, diags = self._run(
            extra_changes=[chg(3, "2026-10-14", "entity", "default|u", "资产字段", "ADD")],
            extra_entity=entity_extra)
        assert any(d["code"] == "new_source_unreferenced" and "'u'" in d["message"] for d in diags)


class TestScenarioModifyType:
    """场景 B（第二刀）：目标字段类型修改（属性级 MODIFY·目标字段类型）——
    三方对账（变更前≈baseline/变更后≈mapping 当前态）+ 漂移豁免。"""

    def _run(self, modify_col="目标字段类型", before="VARCHAR(64)", after="VARCHAR(128)",
             ttype="VARCHAR(128)", extra_changes=(), drop_order_row=False):
        entity = [ent(), ent(table="dim_channel", alias="c")]
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id",
                      ttype=ttype),
                 attr()]
        if drop_order_row:
            attrs = [attr()]  # 当前态无 order_id 行
        mapping = mapping_of(entity, attrs)
        all_changes = [chg(1, "2026-10-14", "attr", "default|order_id", modify_col, "MODIFY",
                           before, after)] + list(extra_changes)
        batch = [_to_entry(c) for c in all_changes]
        renames, rest = synth_renames(batch)
        return extract_and_check(rest, renames, mapping, facts_fixture(), "202610", "")

    def test_valid_modify_collected(self):
        add, ns, mods, mlogics, uns, diags = self._run()
        assert mods and mods[0]["field"] == "order_id"
        assert mods[0]["old_type"] == "VARCHAR(64)" and mods[0]["new_type"] == "VARCHAR(128)"
        assert not [d for d in diags if d["level"] == "error"], diags

    def test_before_mismatch_blocks(self):
        add, ns, mods, mlogics, uns, diags = self._run(before="numeric(10,2)")  # baseline 是 varchar
        assert any(d["code"] == "modify_before_mismatch" for d in diags if d["level"] == "error")

    def test_after_mismatch_blocks(self):
        add, ns, mods, mlogics, uns, diags = self._run(after="VARCHAR(256)")  # 当前态是 128
        assert any(d["code"] == "modify_after_mismatch" for d in diags if d["level"] == "error")

    def test_not_in_baseline_blocks(self):
        add, ns, mods, mlogics, uns, diags = self._run(after="", ttype="VARCHAR(128)")
        # channel_name 不在 baseline——用它测 not_in_baseline
        entity = [ent(), ent(table="dim_channel", alias="c")]
        mapping = mapping_of(entity, [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id"), attr(ttype="VARCHAR(128)")])
        batch = [_to_entry(chg(1, "2026-10-14", "attr", "default|channel_name",
                               "目标字段类型", "MODIFY", "VARCHAR(64)", "VARCHAR(128)"))]
        renames, rest = synth_renames(batch)
        _, _, _, _, _, diags = extract_and_check(rest, renames, mapping, facts_fixture(), "202610", "")
        assert any(d["code"] == "modify_field_not_in_baseline" for d in diags if d["level"] == "error")

    def test_modify_add_conflict_blocks(self):
        add, ns, mods, mlogics, uns, diags = self._run(
            extra_changes=[chg(2, "2026-10-14", "attr", "default|order_id", "资产字段", "ADD")])
        assert any(d["code"] == "modify_add_conflict" for d in diags if d["level"] == "error")

    def test_drift_exempted_for_declared_modify(self):
        """漂移豁免：申报 MODIFY 的 order_id 类型 diff 不再报漂移 warn；
        未申报的存量字段（cust_id）改类型照报——引导走正路。"""
        entity = [ent(), ent(table="dim_channel", alias="c")]
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di", scol="order_id",
                      ttype="VARCHAR(128)"),
                 attr(tcol="cust_id", alias="a", stable="ods_trade_order_di", scol="cust_id",
                      ttype="VARCHAR(200)"),  # 未申报——漂移照报
                 attr()]
        mapping = mapping_of(entity, attrs)
        batch = [_to_entry(chg(1, "2026-10-14", "attr", "default|order_id",
                               "目标字段类型", "MODIFY", "VARCHAR(64)", "VARCHAR(128)"))]
        renames, rest = synth_renames(batch)
        _, _, _, _, _, diags = extract_and_check(rest, renames, mapping, facts_fixture(), "202610", "")
        drifts = [d for d in diags if d["code"] == "field_type_drift"]
        assert drifts and all("cust_id" in d["message"] for d in drifts), \
            f"order_id 应豁免、cust_id 应报: {drifts}"

    def test_other_column_modify_still_unsupported(self):
        """映射规则列的 MODIFY（分类变更）仍待扩展——映射表达式列已被第三刀消费。"""
        add, ns, mods, mlogics, uns, diags = self._run(modify_col="映射规则")
        assert not mods and not mlogics
        assert any(u["change_type"] == "modify" for u in uns)


class TestScenarioModifyLogic:
    """第三刀（2026-10-22）：口径修改（属性级 MODIFY·映射表达式）——对账分级
    （变更后≈当前态 error / 变更前≈baseline 宽松 warn）→ modify_logic_fields 段。"""

    def _run(self, column="映射表达式", before="", after="nvl(a.amt,0)",
             expr="nvl(a.amt,0)", key="order_id", with_baseline_logic=False,
             extra_attr=None):
        entity = [ent(), ent(table="dim_channel", alias="c")]
        attrs = [attr(tcol="order_id", alias="a", stable="ods_trade_order_di",
                      scol="order_id", ttype="VARCHAR(64)"),
                 attr()]
        attrs[0]["映射表达式"] = expr  # 当前态终态口径
        if extra_attr:
            attrs.append(extra_attr)
        f = facts_fixture()
        if with_baseline_logic:  # 注入 baseline 旧口径（demo 桶为空——手工造落点）
            f["logics"]["order_id"] = [{"rule": "R0002", "kind": "processed",
                                        "logic": "a.amt"}]
        mapping = mapping_of(entity, attrs)
        all_changes = [chg(1, "2026-10-14", "attr", f"default|{key}", column, "MODIFY",
                           before, after)]
        batch = [_to_entry(c) for c in all_changes]
        renames, rest = synth_renames(batch)
        return extract_and_check(rest, renames, mapping, f, "202610", "")

    def test_valid_modify_logic_collected(self):
        add, ns, mods, mlogics, uns, diags = self._run(before="a.amt")
        assert mlogics and mlogics[0]["field"] == "order_id"
        assert mlogics[0]["new_expr"] == "nvl(a.amt,0)"
        assert not [d for d in diags if d["level"] == "error"], diags

    def test_after_mismatch_blocks(self):
        add, ns, mods, mlogics, uns, diags = self._run(after="wrong_expr")
        assert any(d["code"] == "modify_after_mismatch" for d in diags if d["level"] == "error")
        assert not mlogics

    def test_before_mismatch_warns_not_blocks(self):
        """变更前与 baseline 旧口径不一致 → warn 宽松（baseline 可能存翻译态）。"""
        add, ns, mods, mlogics, uns, diags = self._run(
            before="完全不相关的旧口径", with_baseline_logic=True)
        assert mlogics, "宽松 warn 不应丢弃条目"
        assert any(d["code"] == "modify_expr_before_mismatch"
                   for d in diags if d["level"] == "warn")

    def test_not_in_baseline_blocks(self):
        """当前态有行但不在 baseline 存量 → error（口径修改只针对存量）。"""
        add, ns, mods, mlogics, uns, diags = self._run(
            key="ghost_f", after="", expr="",
            extra_attr=attr(tcol="ghost_f", alias="a", stable="ods_trade_order_di",
                            scol="x"))
        assert any(d["code"] == "modify_field_not_in_baseline" for d in diags if d["level"] == "error")

    def test_rule_column_still_unsupported(self):
        add, ns, mods, mlogics, uns, diags = self._run(column="映射规则")
        assert not mlogics
        assert any(u["change_type"] == "modify" for u in uns)
