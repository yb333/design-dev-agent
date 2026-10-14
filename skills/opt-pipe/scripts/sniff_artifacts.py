"""sniff_artifacts —— 制品目录 sheet 指纹判型（opt-pipe 步骤 0 的 artifacts 目录通道）。

用户视角：发起优化只给一个制品目录（文件命名/形态零要求）——本脚本按平台 excel 的
固定 sheet 结构判型，产出 manifest 的 artifacts 形态（机器格式），engineer 直接拼
manifest 起 dws-reverser。判不了/多义 → 原样上抛由 engineer question 人（一次性指认）。

指纹（2026-10-14 实测三类完全不重叠，对齐 analyzer 解析面认知）：
  术加 = 含 RULE 且 含 TargetFields；LTS = 含 tasks/jobs/taskParams；DQ = 含 规则基本信息；
  DDL = .sql 文件（后缀即判）。命中多个型 = 多义；一个不中 = unknown。

生命周期：逆向侧 15 号（manifest 装配层自动分拣）实施后本脚本退役（判型下沉 analyzer，
知识单源归解析侧）——见 docs/specs/opt/12-逆向侧需求 §三。
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

# 型 → 判定 sheet 集合（全命中才算该型；集合间无交集）
FINGERPRINTS = {
    "execution_tasks": {"RULE", "TargetFields"},
    "lts": {"tasks", "jobs", "taskParams"},
    "dq": {"规则基本信息"},
}


def sniff_dir(dir_path: Path) -> dict:
    """判型目录 → {型: [文件], ddl: [文件], unknown: [文件], ambiguous: [文件]}。"""
    out = {k: [] for k in list(FINGERPRINTS) + ["ddl", "unknown", "ambiguous"]}
    if not dir_path.is_dir():
        print(f"SNIFF_ERROR: 不是目录: {dir_path}", file=sys.stderr)
        raise SystemExit(2)
    for p in sorted(dir_path.rglob("*")):
        if not p.is_file() or p.name.startswith((".", "~$")):
            continue
        rel = str(p.relative_to(dir_path))
        if p.suffix.lower() == ".sql":
            out["ddl"].append(rel)
            continue
        if p.suffix.lower() not in (".xlsx", ".xls", ".xlsm"):
            out["unknown"].append(rel)
            continue
        try:
            wb = load_workbook(p, read_only=True)
            sheets = set(wb.sheetnames)
            wb.close()
        except Exception as e:
            out["unknown"].append(f"{rel}（读取失败: {e}）")
            continue
        hits = [k for k, need in FINGERPRINTS.items() if need <= sheets]
        if len(hits) == 1:
            out[hits[0]].append(rel)
        elif len(hits) > 1:
            out["ambiguous"].append(f"{rel}（同时命中 {hits}）")
        else:
            out["unknown"].append(f"{rel}（sheet={sorted(sheets)[:5]}）")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="制品目录 sheet 指纹判型（目录 → manifest artifacts 形态）")
    ap.add_argument("--dir", required=True, help="制品目录路径（用户发起时给的目录）")
    args = ap.parse_args(argv)
    result = sniff_dir(Path(args.dir))
    # 缺省语义：目录里没有 DQ 型文件 = 按声明无处理（declared_absent）——三态的"沉默"
    # 由 engineer 在闸口①'披露纠错（"按无 DQ 处理入档，如有遗漏请说明"），不前置停等。
    manifest_arts = {
        "execution_tasks": result["execution_tasks"],
        "lts": result["lts"],
        "dq": ({"status": "provided", "paths": result["dq"]} if result["dq"]
               else {"status": "declared_absent"}),
        "ddl": result["ddl"],
    }
    print(json.dumps({"manifest_artifacts": manifest_arts,
                      "unknown": result["unknown"], "ambiguous": result["ambiguous"],
                      "notes": "多文件/多义/无法识别由 engineer question 一次性指认"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
