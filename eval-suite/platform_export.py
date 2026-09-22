#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""内网评测平台导入集一键构造器（2026-09-22）。

本地试点 → 平台口径确认 → 内网真实案例，同一条命令：

  # 本地试点（eval-suite 虚拟案例 × 10_project_deliver 产物档案 join）
  python3 platform_export.py --source evalsuite
  # 内网真实案例（10_project_deliver 档案直读，兼容老式平铺与 build/ 新布局）
  python3 platform_export.py --source archive --root <内网10_project_deliver路径>

产物三件（过程可视可回溯）：
  out/platform_import.xlsx            平台导入文件
  out/platform_import.csv             同内容 csv——本地 deepeval 预检（EvaluationDataset 加载）
  out/platform_import.manifest.json   每案例取料清单（用了哪个档案/各列字符数/跳过原因）

数据语义：input / expected_output = 案例设计侧构造（任务指令 + 完成标准）；
actual_output = 真实运行档案摘录（ts.md §1 概述 + 产物清单），不编造输出。
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import datetime
from pathlib import Path

# ── 平台口径（唯一调整点：平台确认表头全集/指标合法值后只改这里）─────────
COLUMNS = [
    "metric_name",
    "input",
    "actual_output",
    "expected_output",
    "trace",
    "turns",
    "tools_called",
    "expected_tools",
]
DEFAULT_METRIC_NAME = "任务成功率"  # ← 换成平台评估器的准确名称
EMPTY_COLUMNS = ["trace", "turns", "tools_called", "expected_tools"]  # 任务成功率评估器不消费，留空

CAP_INPUT = 1200
CAP_ACTUAL = 2000
CAP_EXPECTED = 900
CAP_RS_DIGEST = 500


def clip(text: str, cap: int) -> str:
    if len(text) <= cap:
        return text
    return text[:cap] + "…（截断）"


def load_json(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


# ── 档案取料（老式平铺 / build 新布局 / opt 现场，确定性候选路径）─────────

def find_artifact_root(case_dir: Path):
    """case_dir/ddlc_design_dev → 产物根（build/ 优先，ts.md 在位才算数）。"""
    dd = case_dir / "ddlc_design_dev"
    if not dd.is_dir():
        return None
    build = dd / "build"
    if (build / "ts.md").is_file():
        return build
    if (dd / "ts.md").is_file():
        return dd
    return None


def find_rs_input(case_dir: Path):
    dd = case_dir / "ddlc_design_dev"
    candidates = [
        dd / "_internal" / "rs_input.json",
        dd / "build" / "_internal" / "rs_input.json",
    ]
    if dd.is_dir():
        for child in sorted(dd.iterdir()):
            if child.name.startswith("opt_") and child.is_dir():
                candidates.append(child / "rs_input.json")
    for c in candidates:
        if c.is_file():
            return c
    return None


def ts_md_overview(art_root: Path) -> str:
    """ts.md §1 概述段切片（各版本 ts.md 都有，确定性摘要源）。"""
    try:
        text = (art_root / "ts.md").read_text(encoding="utf-8")
    except Exception:
        return ""
    start = text.find("\n## 1.")
    if start < 0:
        return clip(text[:800], 800)
    end = text.find("\n## ", start + 5)
    seg = text[start + 1 : end if end > 0 else start + 1600]
    return seg.strip()


def list_artifacts(art_root: Path) -> list[str]:
    """产物清单：ddl/ + etl/（老档案 select/ 兜底），规则文件名有序。"""
    out = []
    for sub in ("ddl", "etl", "select"):
        d = art_root / sub
        if d.is_dir():
            for f in sorted(d.iterdir()):
                if f.suffix == ".sql":
                    out.append(f"{sub}/{f.name}")
    return out


# ── 行构造 ────────────────────────────────────────────────

def build_actual(art_root: Path, provenance: str) -> str:
    overview = ts_md_overview(art_root)
    files = list_artifacts(art_root)
    parts = ["设计摘要（ts.md §1 概述，真实运行档案摘录）：", overview, "", "产物清单："]
    parts += [f"- {f}" for f in files]
    parts += ["", f"档案来源：{provenance}"]
    return "\n".join(parts)


def row_from_archive(case_dir: Path, metric: str):
    """10_project_deliver 案例目录 → 一行（task 粒度）。"""
    art_root = find_artifact_root(case_dir)
    if art_root is None:
        return None, "无 ts.md 产物档案（未完成或目录不合规）"

    rs_path = find_rs_input(case_dir)
    rs = load_json(rs_path) if rs_path else None
    if rs is None:
        return None, "rs_input.json 缺失（无法派生 input/expected）"

    meta = rs.get("meta", {}) or {}
    target = meta.get("target", {}) or {}
    schema, table = target.get("schema", ""), target.get("table", "")
    full = f"{schema}.{table}" if schema else table
    grain = meta.get("grain", "")
    desc = clip(target.get("description", "") or target.get("cn", ""), 200)
    fields = rs.get("field_mappings", []) or []
    sources = rs.get("source_tables", []) or []
    sched = rs.get("schedule", {}) or {}
    dq = rs.get("dq_requirements", []) or []

    def src_name(s):
        if isinstance(s, dict):
            return s.get("table") or s.get("name") or str(s)
        return str(s)

    inp = (
        f"任务：为 {full} 完成设计、评审、开发及测试（输入=RS+mapping，流程=设计→闸口→编码→UT）。\n"
        f"表说明：{desc}\n粒度：{grain}\n调度：{sched.get('strategy', '')} / {sched.get('frequency', '')}\n"
        f"映射字段 {len(fields)} 项，来源表 {len(sources)} 张。"
    )
    expected = (
        "任务完成标准（自 RS 派生）：\n"
        f"- 产出 TS 设计（ts.md/ts.json），目标表 {full}，粒度：{grain}\n"
        f"- 覆盖全部 {len(fields)} 个映射字段的加工口径，无遗漏无幻觉字段\n"
        f"- 建表 DDL（create_table_{table}.sql）与结构一致；ETL 每规则一个 SELECT 文件\n"
        f"- 调度方案：{sched.get('strategy', '')}"
        + (f"；DQ 检查 {len(dq)} 条落地为 dq/*.sql" if dq else "")
    )
    actual = build_actual(art_root, str(case_dir))
    row = {c: "" for c in COLUMNS}
    row["metric_name"] = metric
    row["input"] = clip(inp, CAP_INPUT)
    row["actual_output"] = clip(actual, CAP_ACTUAL)
    row["expected_output"] = clip(expected, CAP_EXPECTED)
    return row, None


def asset_name(dirname: str) -> str:
    prefix = dirname.split("_", 1)
    return prefix[1] if len(prefix) == 2 and prefix[0].isdigit() else dirname


def match_deliver(asset: str, deliver_root: Path):
    """按资产名 join 10_project_deliver（多档取 ts.md 在位且 mtime 最新）。"""
    if not deliver_root.is_dir():
        return None
    cands = [
        d
        for d in deliver_root.iterdir()
        if d.is_dir() and (d.name == asset or d.name.endswith("_" + asset) or d.name.startswith(asset + "_"))
    ]
    cands = [d for d in cands if find_artifact_root(d) is not None]
    if not cands:
        return None
    return max(cands, key=lambda d: (d / "ddlc_design_dev").stat().st_mtime)


def row_from_evalsuite(case_dir: Path, deliver_root: Path, metric: str):
    """eval-suite 案例 → 一行：input/expected 来自 expectations.json，actual 来自产物档案。"""
    exp = load_json(case_dir / "expectations.json")
    if exp is None:
        return None, "expectations.json 缺失"

    prompt = exp.get("prompt", "")
    target = exp.get("target_table", "")
    rs_digest = ""
    rs_md = case_dir / "RS.md"
    if rs_md.is_file():
        try:
            rs_digest = clip(rs_md.read_text(encoding="utf-8").strip(), CAP_RS_DIGEST)
        except Exception:
            rs_digest = ""
    inp = f"任务：{prompt}\n目标表：{target}\n需求摘要（RS）：\n{rs_digest}"

    scoring = exp.get("scoring", {}) or {}
    checks = exp.get("checks", []) or []
    files = [p for c in checks if c.get("type") == "files_exist" for p in c.get("paths", [])]
    exp_lines = ["任务完成标准："]
    if scoring:
        exp_lines.append("- 评分维度：" + " / ".join(f"{k}{v}" for k, v in scoring.items()))
    if files:
        exp_lines.append("- 应产出文件：" + ", ".join(files[:12]))
        if len(files) > 12:
            exp_lines.append(f"  （等共 {len(files)} 个）")
    expected = "\n".join(exp_lines)

    deliver = match_deliver(asset_name(case_dir.name), deliver_root)
    if deliver is None:
        return None, f"10_project_deliver 无匹配产物档案（资产 {asset_name(case_dir.name)}）"
    art_root = find_artifact_root(deliver)
    actual = build_actual(art_root, str(deliver))
    row = {c: "" for c in COLUMNS}
    row["metric_name"] = metric
    row["input"] = clip(inp, CAP_INPUT)
    row["actual_output"] = clip(actual, CAP_ACTUAL)
    row["expected_output"] = clip(expected, CAP_EXPECTED)
    return row, None


# ── 输出 ─────────────────────────────────────────────────

def write_outputs(rows, manifest, out_base: Path):
    out_base.parent.mkdir(parents=True, exist_ok=True)

    csv_path = out_base.with_suffix(".csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)

    xlsx_path = out_base.with_suffix(".xlsx")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "eval_import"
        ws.append(COLUMNS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r[c] for c in COLUMNS])
        widths = {"metric_name": 14, "input": 60, "actual_output": 80, "expected_output": 60}
        for i, col in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
        ws.freeze_panes = "A2"
        wb.save(xlsx_path)
    except ImportError:
        print("⚠ openpyxl 不可用，只产出 csv（平台如只收 xlsx 请先 pip install openpyxl）", file=sys.stderr)
        xlsx_path = None

    manifest_path = out_base.with_name(out_base.name + ".manifest.json")
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return csv_path, xlsx_path, manifest_path


def main():
    ap = argparse.ArgumentParser(description="内网评测平台导入集一键构造器")
    ap.add_argument("--source", choices=["evalsuite", "archive"], default="evalsuite",
                    help="evalsuite=本地试点（cases × 产物档案 join）/ archive=档案直读（内网真实案例）")
    ap.add_argument("--root", default=None,
                    help="案例根目录：evalsuite 默认本脚本所在 eval-suite；archive 必填（10_project_deliver）")
    ap.add_argument("--deliver-root", default=None,
                    help="evalsuite 源的产物档案根（默认兄弟目录 10_project_deliver）")
    ap.add_argument("--out", default=None, help="输出基准路径（默认 out/platform_import[[_source]].xlsx）")
    ap.add_argument("--metric-name", default=DEFAULT_METRIC_NAME, help="平台评估器名称")
    args = ap.parse_args()

    here = Path(__file__).resolve().parent
    if args.source == "evalsuite":
        root = Path(args.root) if args.root else here / "cases"
        deliver_root = Path(args.deliver_root) if args.deliver_root else here.parent / "10_project_deliver"
        iter_case = lambda: sorted(p for p in root.iterdir() if p.is_dir())
        make_row = lambda d: row_from_evalsuite(d, deliver_root, args.metric_name)
    else:
        if not args.root:
            ap.error("--source archive 需要 --root（10_project_deliver 路径）")
        root = Path(args.root)
        deliver_root = None
        iter_case = lambda: sorted(p for p in root.iterdir() if p.is_dir())
        make_row = lambda d: row_from_archive(d, args.metric_name)

    rows, manifest = [], {"generated_at": datetime.now().isoformat(timespec="seconds"),
                          "metric_name": args.metric_name, "source": args.source, "cases": []}
    for case_dir in iter_case():
        if case_dir.name.startswith("_") or case_dir.name.startswith("."):
            continue
        row, skip = make_row(case_dir)
        entry = {"case": case_dir.name}
        if row:
            entry.update({
                "included": True,
                "len_input": len(row["input"]),
                "len_actual": len(row["actual_output"]),
                "len_expected": len(row["expected_output"]),
                "empty_columns": EMPTY_COLUMNS,
            })
            rows.append(row)
        else:
            entry.update({"included": False, "skip_reason": skip})
        manifest["cases"].append(entry)

    out_base = Path(args.out) if args.out else here / "out" / (
        "platform_import" if args.source == "evalsuite" else "platform_import_intranet")
    csv_path, xlsx_path, manifest_path = write_outputs(rows, manifest, out_base)

    print(f"来源={args.source}  案例目录={root}  入集={len(rows)} 行  指标={args.metric_name}")
    for e in manifest["cases"]:
        mark = "✓" if e["included"] else "✗"
        info = f"in={e['len_input']} actual={e['len_actual']} exp={e['len_expected']}" if e["included"] else e["skip_reason"]
        print(f"  {mark} {e['case']:<32} {info}")
    print(f"产物：{xlsx_path or '（xlsx 跳过）'} / {csv_path} / {manifest_path}")


if __name__ == "__main__":
    main()
